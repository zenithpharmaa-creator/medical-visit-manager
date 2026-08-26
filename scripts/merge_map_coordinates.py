import json
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[1]
SOURCE = Path("/home/ubuntu/upload/Map.xlsx")
TARGET = PROJECT / "data" / "embedded-entities.ts"
REPORT = PROJECT / "data" / "map-merge-report.json"


def clean(value):
    if value is None:
        return ""
    text = str(value).strip()
    text = re.sub(r"[\u064B-\u065F\u0670]", "", text)
    text = text.replace("ـ", "")
    text = re.sub(r"\s+", " ", text)
    return text


def key(value):
    return re.sub(r"[^\w\u0600-\u06FF]+", "", clean(value).casefold())


def number(value):
    if value is None or value == "":
        return None
    text = clean(value).replace("٫", ".").replace("٬", "")
    text = text.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789"))
    try:
        result = float(text)
    except ValueError:
        return None
    return result if -180 <= result <= 180 else None


def region(value):
    text = clean(value)
    if "شمال" in text:
        return "north"
    if "جنوب" in text:
        return "south"
    if "صنعاء" in text:
        return "sanaa"
    return None


def kind(value):
    text = clean(value)
    mapping = {
        "عيادة": "عيادة",
        "مستشفى حكومي": "مستشفى حكومي",
        "مستشفى خاص": "مستشفى خاص",
        "مركز حكومي": "مركز حكومي",
        "مركز خاص": "مركز خاص",
        "برج طبي": "برج طبي",
        "صيدلية": "صيدلية",
        "مستشفى": "مستشفى",
        "مركز طبي": "مركز طبي",
    }
    return mapping.get(text, text or "عيادة")


def entity_match(entity, row):
    name_match = key(entity.get("name")) == key(row.get("name"))
    if not name_match:
        return False
    row_address = key(row.get("address"))
    entity_address = key(entity.get("address"))
    if row_address and entity_address and row_address != entity_address:
        return False
    return True


raw = TARGET.read_text(encoding="utf-8")
start = raw.index("= [") + 2
end = raw.rindex("];" ) + 1
entities = json.loads(raw[start:end])
original_count = len(entities)

wb = load_workbook(SOURCE, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
headers = [clean(value) for value in rows[0]]
records = []
for values in rows[1:]:
    row = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
    name = clean(row.get("اسم الطبيب"))
    if name:
        records.append({
            "name": name,
            "kind": kind(row.get("نوع الجهة")),
            "specialty": clean(row.get("التخصص")),
            "category": clean(row.get("الفئة")),
            "region": region(row.get("منطقة الزيارة")),
            "geographicLocation": clean(row.get("الموقع الجغرافي")),
            "latitude": number(row.get("خط العرض")),
            "longitude": number(row.get("عمود1")),
            "nearestPharmacy": clean(row.get("اقرب صيدلية")),
            "address": clean(row.get("العنوان")),
            "notes": clean(row.get("الملاحضات")),
        })

matched = 0
updated_coordinates = 0
updated_fields = 0
added = 0
skipped_duplicate_rows = 0
unmatched = []
seen_source_keys = set()

for row in records:
    source_key = (key(row["name"]), key(row["address"]), key(row["specialty"]), key(row["category"]))
    if source_key in seen_source_keys:
        skipped_duplicate_rows += 1
        continue
    seen_source_keys.add(source_key)
    index = next((i for i, entity in enumerate(entities) if entity_match(entity, row)), None)
    if index is None:
        entity_id = f"excel-entity-{len(entities) + 1}"
        entity = {
            "id": entity_id,
            "name": row["name"],
            "kind": row["kind"],
            "createdAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        }
        for field in ("specialty", "category", "region", "geographicLocation", "nearestPharmacy", "address", "notes"):
            if row[field]:
                entity[field] = row[field]
        if row["latitude"] is not None and row["longitude"] is not None:
            entity["geographicCoordinates"] = {"latitude": row["latitude"], "longitude": row["longitude"], "accuracy": None}
        entities.append(entity)
        added += 1
        continue

    matched += 1
    entity = entities[index]
    changed = False
    for field in ("kind", "specialty", "category", "region", "geographicLocation", "nearestPharmacy", "address", "notes"):
        if not entity.get(field) and row[field]:
            entity[field] = row[field]
            changed = True
            updated_fields += 1
    if row["latitude"] is not None and row["longitude"] is not None:
        existing = entity.get("geographicCoordinates") or {}
        if not isinstance(existing, dict) or existing.get("latitude") is None or existing.get("longitude") is None:
            entity["geographicCoordinates"] = {"latitude": row["latitude"], "longitude": row["longitude"], "accuracy": existing.get("accuracy") if isinstance(existing, dict) else None}
            updated_coordinates += 1
            changed = True

# Final defensive deduplication: retain the most complete first record for each identity.
deduped = []
identity_index = {}
removed_duplicates = 0
for entity in entities:
    identity = (key(entity.get("name")), key(entity.get("address")), key(entity.get("specialty")), key(entity.get("category")))
    if identity in identity_index:
        prior = deduped[identity_index[identity]]
        for field, value in entity.items():
            if field not in prior or prior.get(field) in (None, ""):
                prior[field] = value
        removed_duplicates += 1
    else:
        identity_index[identity] = len(deduped)
        deduped.append(entity)
entities = deduped

TARGET.write_text(
    "// Generated from Map.xlsx. Coordinates merged without duplicate entities.\n"
    'import type { MedicalEntity } from "@/shared/visit-types";\n'
    "export const EMBEDDED_ENTITIES: MedicalEntity[] = "
    + json.dumps(entities, ensure_ascii=False, indent=2)
    + ";\n",
    encoding="utf-8",
)

with_coordinates = sum(1 for entity in entities if isinstance(entity.get("geographicCoordinates"), dict) and entity["geographicCoordinates"].get("latitude") is not None and entity["geographicCoordinates"].get("longitude") is not None)
report = {
    "source": str(SOURCE),
    "sourceRows": len(records),
    "originalEmbeddedEntities": original_count,
    "finalEmbeddedEntities": len(entities),
    "matchedRows": matched,
    "newEntities": added,
    "coordinatesAdded": updated_coordinates,
    "otherFieldsBackfilled": updated_fields,
    "duplicateSourceRowsSkipped": skipped_duplicate_rows,
    "duplicateEntitiesRemoved": removed_duplicates,
    "entitiesWithCompleteCoordinates": with_coordinates,
    "entitiesWithoutCompleteCoordinates": len(entities) - with_coordinates,
    "generatedAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
}
REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(report, ensure_ascii=False))
